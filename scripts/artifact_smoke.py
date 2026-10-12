from __future__ import annotations

import argparse
import calendar
import io
import os
import socket
import subprocess
import sys
import tempfile
import time
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import httpx
from icalendar import Calendar
from openpyxl import load_workbook


def smoke_payload() -> dict[str, Any]:
    year, month = 2026, 2
    previous = (date(year, month, 1) - timedelta(days=1)).isoformat()
    names = [(f"M{i}", "moscow") for i in range(4)] + [(f"K{i}", "khabarovsk") for i in range(2)]
    return {
        "year": year,
        "month": month,
        "solver": "cpsat",
        "seed": 42,
        "employees": [
            {"name": name, "city": city, "schedule_type": "flexible"} for name, city in names
        ],
        "carry_over": [
            {
                "employee_name": name,
                "source": "new_employee",
                "no_previous_shifts": True,
                "as_of": previous,
            }
            for name, _city in names
        ],
        "production_calendar": {
            "year": year,
            "month": month,
            "source": "delivery_fixture",
            "fetched_at": "2026-10-10T00:00:00Z",
            "verified": True,
            "days": [
                {
                    "date": f"{year}-{month:02d}-{day:02d}",
                    "kind": "off" if calendar.weekday(year, month, day) >= 5 else "working",
                }
                for day in range(1, calendar.monthrange(year, month)[1] + 1)
            ],
        },
    }


def wait_for_url(url: str, timeout: float = 45) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            if httpx.get(url, timeout=2).status_code == 200:
                return
        except httpx.HTTPError:
            pass
        time.sleep(0.25)
    raise RuntimeError(f"Runtime did not become healthy: {url}")


def verify_api(base_url: str) -> None:
    wait_for_url(base_url + "/openapi.json")
    with httpx.Client(base_url=base_url, timeout=120) as client:
        generated = client.post("/api/v1/schedule/generate", json=smoke_payload())
        generated.raise_for_status()
        schedule = generated.json()
        if schedule.get("status") != "ready" or len(schedule["days"]) != 28:
            raise RuntimeError("Installed API did not generate a checked complete month")
        spreadsheet = client.post("/api/v1/export/revision/xls", json=schedule)
        spreadsheet.raise_for_status()
        workbook = load_workbook(io.BytesIO(spreadsheet.content))
        if not workbook.sheetnames:
            raise RuntimeError("Installed API produced an empty workbook")
        personal = client.post(
            "/api/v1/export/revision/ics", params={"employee_name": "K0"}, json=schedule
        )
        personal.raise_for_status()
        events = Calendar.from_ical(personal.content).walk("VEVENT")
        nights = [
            event
            for event in events
            if "Ночная" in str(event.get("summary", "")) or "Ночь" in str(event.get("summary", ""))
        ]
        if not nights or any(
            (event.decoded("dtend") - event.decoded("dtstart")).total_seconds() != 8 * 3600
            for event in nights
        ):
            raise RuntimeError("Installed API night calendars have incorrect durations")
        pdf = client.post("/api/v1/export/revision/pdf", json=schedule)
        pdf.raise_for_status()
        if not pdf.content.startswith(b"%PDF-"):
            raise RuntimeError("Installed API did not export a PDF")
        schedule["days"][0]["night"] = []
        schedule["status"] = "ready"
        blocked = client.post("/api/v1/export/revision/xls", json=schedule)
        if blocked.status_code != 409:
            raise RuntimeError("Installed API accepted a forged ready schedule")


def verify_installed(source_root: Path) -> None:
    import json

    import duty_schedule

    imported = Path(duty_schedule.__file__).resolve()
    if source_root.resolve() in imported.parents:
        raise RuntimeError("Artifact smoke imported repository source instead of installed wheel")
    with tempfile.TemporaryDirectory(prefix="duty-wheel-smoke-") as temporary:
        folder = Path(temporary)
        config = folder / "config.yaml"
        config.write_text(json.dumps(smoke_payload()), encoding="utf-8")
        environment = {key: value for key, value in os.environ.items() if key != "PYTHONPATH"}
        cli = subprocess.run(
            [
                "duty-schedule",
                "generate",
                str(config),
                "--output-dir",
                str(folder),
                "--format",
                "all",
            ],
            cwd=folder,
            env=environment,
            text=True,
            capture_output=True,
            timeout=150,
        )
        if cli.returncode != 0:
            raise RuntimeError(f"Installed CLI failed: {cli.stdout}\n{cli.stderr}")
        spreadsheets = list(folder.glob("*.xlsx"))
        calendars = list(folder.glob("*.ics"))
        if len(spreadsheets) != 1 or len(calendars) < 3:
            raise RuntimeError("Installed CLI did not produce its exports")
        load_workbook(spreadsheets[0]).close()
        for path in calendars:
            Calendar.from_ical(path.read_bytes())
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        environment["DUTY_API_AUTH_ENABLED"] = "false"
        with (folder / "api.log").open("w") as log:
            server = subprocess.Popen(
                [
                    sys.executable,
                    "-m",
                    "uvicorn",
                    "duty_schedule.api:create_app",
                    "--factory",
                    "--host",
                    "127.0.0.1",
                    "--port",
                    str(port),
                ],
                cwd=folder,
                env=environment,
                stdout=log,
                stderr=subprocess.STDOUT,
            )
            try:
                verify_api(f"http://127.0.0.1:{port}")
            finally:
                server.terminate()
                try:
                    server.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    server.kill()
                    server.wait(timeout=5)
    print(f"Installed wheel CLI/API smoke passed: {duty_schedule.__version__}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    args = parser.parse_args()
    verify_installed(args.source_root)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
