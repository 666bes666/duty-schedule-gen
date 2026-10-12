import pytest
from openpyxl import load_workbook
from playwright.sync_api import Page, expect

from duty_schedule.models import Schedule
from duty_schedule.project import add_revision, new_project, serialize_project
from tests.ui.fixtures import journey_config, journey_days
from tests.ui.test_project_journey import _download, _upload


@pytest.mark.ui
def test_summary_and_exports_show_same_quality_window(page: Page, streamlit_url: str, tmp_path):
    page.goto(streamlit_url, wait_until="networkidle")
    config = journey_config()
    project = add_revision(new_project(config), Schedule(config=config, days=journey_days(config)))
    _upload(page, "Открыть рабочий проект (.json)", "quality.json", serialize_project(project))
    page.get_by_role("tab", name="Нагрузка", exact=True).click()
    expect(page.get_by_text("Окно качества: month_only", exact=False)).to_be_visible(timeout=20000)
    path = _download(page, page.get_by_role("button", name="Скачать XLS", exact=True), tmp_path)
    workbook = load_workbook(path)
    assert workbook["Качество"].cell(1, 2).value == "month_only"
    assert workbook["Качество"].cell(1, 4).value == "missing"
