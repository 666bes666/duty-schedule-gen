from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

import pytest
from icalendar import Calendar
from openpyxl import load_workbook
from playwright.sync_api import Locator, Page, expect

from duty_schedule.project import new_project, serialize_project
from duty_schedule.ui.config_io import dump_config_yaml
from tests.ui.fixtures import journey_config


def _upload(page: Page, label: str, name: str, data: bytes, valid: bool = True) -> None:
    page.wait_for_function(
        "document.querySelector('[data-testid=stApp]')?.dataset.testScriptState === 'notRunning'"
    )
    uploader = page.get_by_test_id("stFileUploader").filter(has=page.get_by_text(label, exact=True))
    uploader.locator('input[type="file"]').set_input_files(
        {
            "name": name,
            "mimeType": "application/json" if name.endswith("json") else "text/yaml",
            "buffer": data,
        }
    )
    if valid:
        expect(page.get_by_text(f"Загружено: {name}", exact=True)).to_be_visible(timeout=20000)
    else:
        expect(page.get_by_text(name, exact=True)).to_be_visible(timeout=20000)


def _download(page: Page, button: Locator, directory: Path) -> Path:
    with page.expect_download(timeout=20000) as event:
        button.click()
    download = event.value
    path = directory / download.suggested_filename
    download.save_as(path)
    page.wait_for_function(
        "document.querySelector('[data-testid=stApp]')?.dataset.testScriptState === 'notRunning'"
    )
    return path


def _assign_four_hour_workday(page: Page, number: int) -> None:
    page.get_by_role("tab", name="Редактирование", exact=True).click()
    date_control = (
        page.get_by_test_id("stDateInput")
        .filter(has=page.get_by_text("Дата «Р»", exact=True))
        .locator("input")
    )
    date_control.fill(f"2026-04-{number:02d}")
    date_control.press("Tab")
    page.get_by_label("Индивидуальные рабочие часы", exact=True).fill("4")
    page.get_by_role("button", name="Назначить «Р» и подтвердить часы", exact=True).click()
    expect(
        page.get_by_text("Текущая версия изменена и требует проверки", exact=True)
    ).to_be_visible(timeout=20000)
    page.get_by_role("button", name="Проверить текущую версию", exact=True).click()
    message = (
        "Диагностический черновик: обычный экспорт недоступен"
        if number == 2
        else "Текущая версия прошла обязательные проверки"
    )
    expect(page.get_by_text(message, exact=True)).to_be_visible(timeout=20000)
    from tests.ui.test_product_recovery import wait_saved

    wait_saved(page)


def _active_schedule(data: dict) -> dict:
    revision_id = data["active_revisions"][data["selected_period"]]
    return next(item["schedule"] for item in data["revisions"] if item["id"] == revision_id)


@pytest.mark.ui
def test_project_generation_edit_download_and_reopen(
    page: Page, streamlit_url: str, tmp_path: Path
) -> None:
    page.goto(streamlit_url, wait_until="networkidle")
    _upload(
        page,
        "Загрузить конфиг (.yaml)",
        "offline.yaml",
        dump_config_yaml(journey_config()).encode(),
    )
    page.get_by_role("button", name="Сгенерировать расписание", exact=True).click()
    expect(
        page.get_by_text("Текущая версия прошла обязательные проверки", exact=True)
    ).to_be_visible(timeout=30000)
    _assign_four_hour_workday(page, 2)
    expect(
        page.get_by_text("Диагностический черновик: обычный экспорт недоступен", exact=True)
    ).to_be_visible(timeout=20000)
    expect(page.get_by_role("button", name="Скачать XLS", exact=True)).to_have_count(0)
    _assign_four_hour_workday(page, 10)
    expect(
        page.get_by_text("Текущая версия прошла обязательные проверки", exact=True)
    ).to_be_visible(timeout=20000)

    heading = page.get_by_role("heading", name=re.compile("Апрель 2026 · текущая версия"))
    previous_heading = heading.inner_text()
    page.get_by_role("button", name="Сгенерировать расписание", exact=True).click()
    expect(heading).not_to_have_text(previous_heading, timeout=30000)
    expect(
        page.get_by_text("Текущая версия прошла обязательные проверки", exact=True)
    ).to_be_visible()

    xls_path = _download(page, page.get_by_role("button", name="Скачать XLS", exact=True), tmp_path)
    workbook = load_workbook(xls_path)
    sheet = workbook["График дежурств"]
    employee_row = next(
        row for row in range(3, sheet.max_row + 1) if sheet.cell(row, 1).value == "M0"
    )
    assert sheet.cell(employee_row, 4).value == "День (4ч)"
    assert sheet.cell(employee_row, 12).value == "День (4ч)"

    page.get_by_text("Скачать календарь для сотрудника", exact=True).click()
    ics_path = _download(page, page.get_by_role("button", name="M0", exact=True), tmp_path)
    events = [
        item for item in Calendar.from_ical(ics_path.read_bytes()).walk() if item.name == "VEVENT"
    ]
    workday = next(event for event in events if event["DTSTART"].dt.date() == date(2026, 4, 2))
    assert workday["DTSTART"].dt.hour == 9
    assert "4" in str(workday["DESCRIPTION"])
    assert "DTEND" not in workday and "DURATION" not in workday

    saved = _download(
        page, page.get_by_role("button", name="Сохранить рабочий проект", exact=True), tmp_path
    )
    before = _active_schedule(json.loads(saved.read_bytes()))
    page.reload(wait_until="networkidle")
    _upload(page, "Открыть рабочий проект (.json)", "reopen.json", saved.read_bytes())
    expect(
        page.get_by_text("Текущая версия прошла обязательные проверки", exact=True)
    ).to_be_visible(timeout=20000)
    reopened_dir = tmp_path / "reopened"
    reopened_dir.mkdir()
    reopened = _download(
        page, page.get_by_role("button", name="Сохранить рабочий проект", exact=True), reopened_dir
    )
    after = _active_schedule(json.loads(reopened.read_bytes()))
    assert after["days"] == before["days"]
    assert after["metadata"]["carry_over"] == before["metadata"]["carry_over"]
    assert after["monthly_targets"]["M0"]["selected_hours"] == 136


@pytest.mark.ui
def test_invalid_project_upload_preserves_existing_work(
    page: Page, streamlit_url: str, tmp_path: Path
) -> None:
    page.goto(streamlit_url, wait_until="networkidle")
    content = serialize_project(new_project(journey_config()))
    _upload(page, "Открыть рабочий проект (.json)", "valid.json", content)
    before_path = _download(
        page, page.get_by_role("button", name="Сохранить рабочий проект", exact=True), tmp_path
    )
    before = json.loads(before_path.read_bytes())
    _upload(
        page,
        "Открыть рабочий проект (.json)",
        "broken.json",
        b'{"format_version":999}',
        valid=False,
    )
    expect(page.get_by_text("Некорректный файл рабочего проекта", exact=True)).to_be_visible(
        timeout=15000
    )
    after_path = _download(
        page, page.get_by_role("button", name="Сохранить рабочий проект", exact=True), tmp_path
    )
    assert json.loads(after_path.read_bytes()) == before
