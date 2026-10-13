from datetime import date

import pandas as pd
import pytest
from pydantic import ValidationError

from duty_schedule.models import DaySchedule, Employee, Schedule
from duty_schedule.ui.builders import _edit_df_to_schedule, _schedule_to_edit_df
from duty_schedule.ui.components import _carry_from_frame
from duty_schedule.ui.config_io import parse_config_yaml
from tests.ui.fixtures import journey_config


def test_employee_identity_roundtrips_names_and_rejects_blank():
    config = journey_config()
    for name in ["Иванов, Иван", '"Ольга"', " 李 明 "]:
        config.employees[0].name = name
        schedule = Schedule(
            config=config, days=[DaySchedule(date=date(2026, 4, 1), morning=[name])]
        )
        restored = _edit_df_to_schedule(_schedule_to_edit_df(schedule), schedule)
        assert restored.days[0].morning == [name]
        assert restored.config.employees[0].employee_id == config.employees[0].employee_id
    with pytest.raises(ValidationError):
        Employee(name="  ", city="moscow", schedule_type="flexible")


@pytest.mark.parametrize("empty", [None, float("nan")])
def test_cleared_editor_cell_moves_employee_to_day_off(empty):
    config = journey_config()
    schedule = Schedule(config=config, days=[DaySchedule(date=date(2026, 4, 1), morning=["M0"])])
    frame = _schedule_to_edit_df(schedule)
    frame.at[0, "Утро 08–17"] = empty
    restored = _edit_df_to_schedule(frame, schedule)
    assert restored.days[0].morning == []
    assert restored.days[0].day_off == ["M0"]


def test_unchanged_history_preserves_counters_and_source():
    config = journey_config()
    carry = config.carry_over[0]
    carry.consecutive_off = 4
    carry.source = "xls"
    frame = pd.DataFrame(
        [
            {
                "ID": c.employee_id,
                "История известна": True,
                "Новый сотрудник": False,
                "Последний статус": "Выходной",
                "Рабочая серия": 0,
            }
            for c in config.carry_over
        ]
    )
    assert _carry_from_frame(frame, config, date(2026, 3, 31)) == config.carry_over
    frame.loc[1, "Рабочая серия"] = 2
    frame.loc[1, "Последний статус"] = "Утро"
    with pytest.raises(ValueError, match="счётчик"):
        _carry_from_frame(frame, config, date(2026, 3, 31))


def test_legacy_team_lead_migration_is_visible():
    from duty_schedule.ui.config_io import dump_config_yaml

    raw = dump_config_yaml(journey_config()).replace("name: M0", "name: M0\n  team_lead: true", 1)
    with pytest.warns(UserWarning, match="team_lead"):
        config = parse_config_yaml(raw)
    assert config.employees[0].name == "M0"


def test_unknown_input_field_is_rejected():
    config = journey_config().model_dump()
    config["employees"][0]["typo"] = True
    from duty_schedule.models import Config

    with pytest.raises(ValidationError):
        Config.model_validate(config)


def test_optimizer_without_changelog_reports_changes_and_quality():
    from duty_schedule.scheduler.postprocess.quality import optimization_provenance

    config = journey_config()
    before = Schedule(config=config, days=[DaySchedule(date=date(2026, 4, 1), morning=["M0"])])
    after = before.model_copy(deep=True)
    after.days[0].morning = ["M1"]
    result = optimization_provenance(before, after, "test")
    assert result["status"] == "changes"
    assert result["changed_cells"] == 2
    assert result["before"]["days"] != result["after"]["days"]
    assert len(result["quality_before"]) == len(result["quality_labels"])


def test_xls_preserves_exact_names_and_rejects_missing_roster(tmp_path):
    from io import BytesIO

    from openpyxl import load_workbook

    from duty_schedule.export.xls import export_xls
    from duty_schedule.project import ensure_exportable
    from duty_schedule.xls_import import XlsImportError, parse_carry_over_from_xls
    from tests.ui.fixtures import journey_days

    config = journey_config()
    days = journey_days(config)
    old = config.employees[0].name
    name = ' 李, "明" '
    config.employees[0].name = name
    for pin in config.pins:
        if pin.employee_name == old:
            pin.employee_name = name
    for carry in config.carry_over:
        if carry.employee_name == old:
            carry.employee_name = name
    from duty_schedule.models import ShiftType

    for day in days:
        for shift in ShiftType:
            setattr(
                day,
                shift.value,
                [name if value == old else value for value in getattr(day, shift.value)],
            )
    path = export_xls(ensure_exportable(Schedule(config=config, days=days)), tmp_path)
    assert name in [c.employee_name for c in parse_carry_over_from_xls(path.read_bytes())]
    workbook = load_workbook(path)
    sheet = workbook["График дежурств"]
    sheet.cell(3, 1).value = None
    stream = BytesIO()
    workbook.save(stream)
    with pytest.raises(XlsImportError, match="Строка 3"):
        parse_carry_over_from_xls(stream.getvalue())


def test_editing_one_history_preserves_other_complete_states():
    config = journey_config()
    config.carry_over[0].source = "xls"
    config.carry_over[0].consecutive_off = 4
    frame = pd.DataFrame(
        [
            {
                "ID": c.employee_id,
                "История известна": True,
                "Новый сотрудник": False,
                "Последний статус": "Выходной",
                "Рабочая серия": 0,
                "Серия отдыха": c.consecutive_off,
                "Серия одной смены": 0,
            }
            for c in config.carry_over
        ]
    )
    frame.loc[1, ["Последний статус", "Рабочая серия", "Серия отдыха", "Серия одной смены"]] = [
        "Утро",
        2,
        0,
        2,
    ]
    changed = _carry_from_frame(frame, config, date(2026, 3, 31))
    assert changed[0] == config.carry_over[0]
    assert changed[2:] == config.carry_over[2:]
    assert changed[1].consecutive_same_shift == 2
    assert changed[1].source == "manual"


def test_cli_rejects_mismatched_calendar_with_both_periods(tmp_path):
    import yaml
    from typer.testing import CliRunner

    from duty_schedule.cli import app

    raw = journey_config().model_dump(mode="json")
    raw["month"] = 5
    path = tmp_path / "mismatch.yaml"
    path.write_text(yaml.safe_dump(raw, allow_unicode=True))
    result = CliRunner().invoke(app, ["validate", str(path)])
    assert result.exit_code != 0
    assert "2026-04" in result.output
    assert "2026-05" in result.output


def test_ui_period_change_selects_new_calendar_before_validation():
    from pathlib import Path

    from streamlit.testing.v1 import AppTest

    from duty_schedule.project import new_project

    app = AppTest.from_file(str(Path(__file__).resolve().parents[2] / "app.py"))
    app.session_state["autosave_initialized"] = True
    app.session_state["project"] = new_project(journey_config())
    app.run(timeout=20)
    next(widget for widget in app.selectbox if widget.label == "Месяц").select(5)
    app.run(timeout=20)
    assert not app.exception
    assert app.session_state["project"].config.month == 5
    assert app.session_state["project"].config.production_calendar is None


def test_manual_edit_does_not_relabel_old_optimizer_snapshot():
    from duty_schedule.project import add_revision, apply_assignments, current_schedule, new_project
    from duty_schedule.scheduler.postprocess.quality import optimization_provenance
    from duty_schedule.scheduler.validation import validate_schedule_result
    from tests.ui.fixtures import journey_days

    config = journey_config()
    schedule = validate_schedule_result(Schedule(config=config, days=journey_days(config)))
    schedule.metadata["optimization"] = optimization_provenance(schedule, schedule, "test")
    project = add_revision(new_project(config), schedule)
    before = current_schedule(project)
    days = [day.model_copy(deep=True) for day in before.days]
    days[0].morning, days[0].evening = days[0].evening, days[0].morning
    edited = current_schedule(apply_assignments(project, days))
    assert edited.metadata["optimization"]["after"]["revision_id"] == before.revision_id
    assert edited.metadata["optimization"]["after"]["revision_id"] != edited.revision_id


def test_malformed_optimizer_metadata_cannot_replace_open_project():
    import json

    from duty_schedule.project import add_revision, new_project, serialize_project
    from duty_schedule.ui.components import import_project_into_state
    from tests.ui.fixtures import journey_days

    config = journey_config()
    project = add_revision(new_project(config), Schedule(config=config, days=journey_days(config)))
    raw = json.loads(serialize_project(project))
    raw["revisions"][0]["schedule"]["metadata"]["optimization"] = {"status": "changes"}
    state = {"project": project}
    before = serialize_project(project)
    with pytest.raises(ValueError):
        import_project_into_state(state, json.dumps(raw))
    assert serialize_project(state["project"]) == before


def test_restored_repair_candidate_cannot_trust_ready_status():
    from duty_schedule.project import load_project, new_project, serialize_project

    config = journey_config()
    project = new_project(config)
    project.repair_candidate = Schedule(config=config, days=[], status="ready")
    restored = load_project(serialize_project(project))
    assert restored.repair_candidate.status == "diagnostic"


def test_cli_and_ui_reject_mismatched_calendar_atomically(tmp_path):
    import yaml
    from typer.testing import CliRunner

    from duty_schedule.cli import app
    from duty_schedule.project import new_project
    from duty_schedule.ui.components import import_config_into_state

    config = journey_config()
    raw = config.model_dump(mode="json")
    raw["month"] = 5
    text = yaml.safe_dump(raw, allow_unicode=True)
    path = tmp_path / "mismatch.yaml"
    path.write_text(text)
    assert CliRunner().invoke(app, ["validate", str(path)]).exit_code != 0
    original = new_project(config)
    state = {"project": original}
    with pytest.raises(ValueError, match="2026-04.*2026-05"):
        import_config_into_state(state, text)
    assert state["project"] is original
