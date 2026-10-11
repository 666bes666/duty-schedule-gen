from datetime import date
from fractions import Fraction

import pytest

from duty_schedule.models import CarryOverState, DaySchedule, Schedule
from tests.ui.fixtures import journey_config, journey_days


def test_cumulative_weights_use_eligible_roles_and_selected_norm_each_period():
    from duty_schedule.quality_history import weighted_imbalance

    assert weighted_imbalance({"A": 34, "B": 26}, {"A": 320, "B": 240}) == Fraction(4, 7)
    assert weighted_imbalance({"A": 35, "B": 25}, {"A": 320, "B": 240}) == Fraction(10, 7)


def test_missing_and_stale_history_are_not_zero():
    from duty_schedule.quality_history import quality_report

    config = journey_config()
    report = quality_report(Schedule(config=config, days=journey_days(config)))
    assert report["window"] == "month_only"
    assert report["history"] == "missing"
    config.previous_schedule = Schedule(config=config, days=journey_days(config))
    report = quality_report(Schedule(config=config, days=journey_days(config)))
    assert report["history"] == "stale"


def test_previous_active_revision_changes_quality_source():
    from duty_schedule.quality_history import previous_quality_schedule

    config = journey_config()
    from duty_schedule.scheduler.validation import validate_schedule_result

    previous = validate_schedule_result(
        Schedule(config=config, days=journey_days(config), revision_id="source")
    )
    next_config = config.model_copy(deep=True)
    next_config.month = 5
    next_config.production_calendar = None
    next_config.previous_schedule = previous
    assert previous_quality_schedule(next_config)[0].revision_id == "source"
    previous.status = "diagnostic"
    assert previous_quality_schedule(next_config)[0] is None


def test_month_and_boundary_streak_windows_are_distinct():
    from duty_schedule.stats import build_assignments, compute_stats

    config = journey_config()
    config.carry_over[0] = CarryOverState(
        employee_name="M0",
        employee_id=config.employees[0].employee_id,
        last_shift="morning",
        consecutive_working=2,
        consecutive_same_shift=2,
        as_of=date(2026, 3, 31),
        source="manual",
    )
    schedule = Schedule(
        config=config,
        days=[DaySchedule(date=date(2026, 4, n), morning=["M0"]) for n in range(1, 4)],
    )
    stats = compute_stats(schedule, build_assignments(schedule), 17)[0]
    assert stats.max_streak_work == 3
    assert stats.boundary_max_streak_work == 5


def test_project_invalidates_old_quality_source_and_reload():
    from duty_schedule.project import (
        add_revision,
        current_schedule,
        load_project,
        new_project,
        serialize_project,
    )
    from duty_schedule.quality_history import quality_report
    from tests.integration.test_two_month_quality import next_config

    config = next_config()
    previous = config.previous_schedule
    project = add_revision(new_project(previous.config), previous)
    config.previous_schedule = current_schedule(project)
    project = add_revision(project, Schedule(config=config, days=[]))
    old = previous.model_copy(deep=True)
    old.config.seed += 1
    project = add_revision(project, old)
    project.selected_period = "2026-05"
    for source in (project, load_project(serialize_project(project))):
        report = quality_report(current_schedule(source))
        assert report["history"] == "stale"
        assert report["window"] == "month_only"
        assert report["source_revision_id"] is None


def test_new_and_departed_employee_do_not_invent_previous_counts():
    from duty_schedule.models import Employee
    from duty_schedule.quality_history import balance_pools
    from tests.integration.test_two_month_quality import next_config

    config = next_config()
    retired_id = config.employees[0].employee_id
    newcomer = Employee(name="New", city="moscow", schedule_type="flexible")
    config.employees[0] = newcomer
    pool = balance_pools(config)["evening"]
    assert retired_id in pool["weights"]
    assert retired_id not in pool["current_ids"]
    assert newcomer.employee_id not in pool["previous_counts"]
    assert pool["weights"][newcomer.employee_id] == 168


def test_year_boundary_checks_adjacent_period_without_fabricating_history():
    from duty_schedule.quality_history import previous_quality_schedule

    config = journey_config()
    previous = Schedule(config=config, days=journey_days(config))
    current = config.model_copy(deep=True)
    current.year = 2027
    current.month = 1
    current.production_calendar = None
    current.previous_schedule = previous
    previous.config.year = 2026
    previous.config.month = 12
    previous.config.production_calendar = None
    _, status = previous_quality_schedule(current)
    assert status == "unverified"
    previous.config.month = 11
    assert previous_quality_schedule(current)[1] == "stale"


def test_rebinding_updated_previous_revision_reenables_quality():
    from duty_schedule.project import add_revision, current_schedule, new_project
    from duty_schedule.quality_history import previous_quality_schedule
    from duty_schedule.ui.components import _config_for_period
    from tests.integration.test_two_month_quality import next_config

    config = next_config()
    previous = config.previous_schedule
    project = add_revision(new_project(previous.config), previous)
    config.previous_schedule = current_schedule(project)
    project = add_revision(project, Schedule(config=config, days=[]))
    previous.config.seed += 1
    project = add_revision(project, previous)
    fresh = _config_for_period(project, 2026, 5)
    snapshot, state = previous_quality_schedule(fresh)
    assert state == "verified"
    assert snapshot.revision_id == project.active_revisions["2026-04"]


@pytest.mark.parametrize("quality_state", ["stale", "available"])
def test_whatif_preserves_trusted_quality_context(quality_state):
    from duty_schedule.api.whatif_service import apply_patch
    from duty_schedule.quality_history import previous_quality_schedule
    from tests.integration.test_two_month_quality import next_config

    config = next_config()
    config._quality_state = quality_state
    config._quality_origin = "project"
    changed = apply_patch(config, {})
    assert changed._quality_origin == "project"
    assert previous_quality_schedule(changed)[1] == (
        "stale" if quality_state == "stale" else "verified"
    )
    external = apply_patch(config, {"previous_schedule": config.previous_schedule.model_dump()})
    assert external._quality_origin == "external"


def test_pdf_renderer_receives_month_and_boundary_series(monkeypatch):
    import sys
    from types import SimpleNamespace

    from duty_schedule.export.pdf import generate_schedule_pdf
    from duty_schedule.scheduler.validation import validate_schedule_result
    from duty_schedule.stats import build_assignments, compute_stats

    config = journey_config()
    schedule = validate_schedule_result(Schedule(config=config, days=journey_days(config)))
    assert schedule.status == "ready"
    captured = {}

    class CapturingHTML:
        def __init__(self, *, string, url_fetcher):
            captured["html"] = string

        def write_pdf(self, stream):
            stream.write(b"%PDF-test")

    monkeypatch.setitem(sys.modules, "weasyprint", SimpleNamespace(HTML=CapturingHTML))
    monkeypatch.setattr("duty_schedule.export.pdf._safe_url_fetcher", lambda: None)
    assert generate_schedule_pdf(schedule) == b"%PDF-test"
    html = captured["html"]
    assert "Серия в месяце" in html
    assert "Серия с границей" in html
    stats = compute_stats(schedule, build_assignments(schedule), 21)
    for stat in stats:
        assert f"<td>{stat.max_streak_work}</td><td>{stat.boundary_max_streak_work}</td>" in html


def test_xls_external_quality_source_id_is_plain_text(tmp_path):
    from openpyxl import load_workbook

    from duty_schedule.export.xls import export_xls
    from duty_schedule.scheduler import generate_schedule
    from tests.integration.test_two_month_quality import next_config

    config = next_config()
    config.previous_schedule.revision_id = "=1+1"
    config.solver = "cpsat"
    schedule = generate_schedule(config)
    assert schedule.status == "ready"
    workbook = load_workbook(export_xls(schedule, tmp_path))
    cell = workbook["Качество"].cell(2, 3)
    assert cell.data_type == "s"
    assert cell.value == "'=1+1"


@pytest.mark.parametrize("changed", [False, True])
def test_config_save_cannot_resurrect_stale_quality_source(changed):
    from duty_schedule.models import Config
    from duty_schedule.project import add_revision, current_schedule, new_project, update_config
    from duty_schedule.quality_history import quality_report
    from tests.integration.test_two_month_quality import next_config

    config = next_config()
    previous = config.previous_schedule
    project = add_revision(new_project(previous.config), previous)
    config.previous_schedule = current_schedule(project)
    project = add_revision(project, Schedule(config=config, days=[]))
    if changed:
        previous.config.seed += 1
        project = add_revision(project, previous)
    project.selected_period = "2026-05"
    saved = Config.model_validate(current_schedule(project).config.model_dump())
    project = update_config(project, saved)
    schedule = current_schedule(project)
    report = quality_report(schedule)
    assert report["history"] == ("stale" if changed else "verified")
    if changed:
        assert any(d.code == "previous_revision_changed" for d in schedule.diagnostics)
    else:
        assert report["source"] == "project"
